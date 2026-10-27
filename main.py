import flet as ft
from openpyxl import load_workbook

def main(page: ft.Page):
    page.title = "奶蛙抽人"  # 这里的名字可以随便改
    page.theme_mode = ft.ThemeMode.LIGHT
    page.padding = 30

    names = []
    
    # 界面元素
    title_text = ft.Text("奶蛙抽人系统", size=28, weight=ft.FontWeight.BOLD)
    status_text = ft.Text("请先导入 Excel 名单", color="grey")
    result_text = ft.Text("", size=40, weight=ft.FontWeight.BOLD, color="blue", text_align=ft.TextAlign.CENTER)
    
    # 文件选择器
    file_picker = ft.FilePicker()
    
    # 读取Excel逻辑
    async def import_excel(e):
        nonlocal names
        try:
            files = await file_picker.pick_files(
                dialog_title="请选择Excel名单",
                allow_multiple=False,
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx"]
            )
            if files:
                path = files[0].path
                wb = load_workbook(path, read_only=True, data_only=True)
                ws = wb.active
                names = []
                for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                    if row[0] is not None:
                        name = str(row[0]).strip()
                        if name and name not in ["姓名", "名字", "名单"]: # 自动跳过表头
                            names.append(name)
                wb.close()
                
                status_text.value = f"✅ 成功导入 {len(names)} 人"
                status_text.color = "green"
                result_text.value = "准备就绪"
                page.update()
        except Exception as ex:
            status_text.value = f"❌ 导入失败: {str(ex)}"
            status_text.color = "red"
            page.update()

    # 抽人逻辑
    def draw(e):
        import random
        if not names:
            status_text.value = "❌ 请先导入名单！"
            status_text.color = "red"
            page.update()
            return
        
        picked = random.choice(names)
        result_text.value = picked
        status_text.value = f"已从 {len(names)} 人中抽取"
        status_text.color = "grey"
        page.update()

    # 把文件选择器加入页面
    if hasattr(page, "services"):
        page.services.append(file_picker)
    else:
        page.overlay.append(file_picker)

    # 排版
    page.add(
        ft.Column(
            [
                title_text,
                ft.Button("📂 导入 Excel 名单", on_click=import_excel, width=200),
                status_text,
                ft.Divider(height=50),
                result_text,
                ft.Divider(height=50),
                ft.Button("🎯 抽 人", on_click=draw, width=200, height=60),
            ],
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            spacing=20
        )
    )

ft.run(main)
